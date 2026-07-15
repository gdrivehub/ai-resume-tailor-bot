"""
Generates a clean, ATS-friendly DOCX resume from the structured JSON
returned by the AI engine (see TAILOR_RESUME_PROMPT schema).
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ACCENT_COLOR = RGBColor(0x1F, 0x4E, 0x79)


def _add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = ACCENT_COLOR
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    # Bottom border for section heading
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def build_resume_docx(data: dict, output_path: str) -> str:
    doc = Document()

    # Base font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
        section.left_margin = Pt(50)
        section.right_margin = Pt(50)

    # Name header
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(data.get("full_name", "Your Name"))
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = ACCENT_COLOR

    if data.get("contact_line"):
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_run = contact_p.add_run(data["contact_line"])
        contact_run.font.size = Pt(9.5)

    # Summary
    if data.get("summary"):
        _add_heading(doc, "Professional Summary")
        doc.add_paragraph(data["summary"])

    # Skills
    if data.get("skills"):
        _add_heading(doc, "Skills")
        doc.add_paragraph(" • ".join(data["skills"]))

    # Experience
    if data.get("experience"):
        _add_heading(doc, "Experience")
        for job in data["experience"]:
            p = doc.add_paragraph()
            title_run = p.add_run(f"{job.get('title', '')} — {job.get('company', '')}")
            title_run.bold = True
            if job.get("dates") or job.get("location"):
                meta_p = doc.add_paragraph()
                meta_run = meta_p.add_run(
                    " | ".join(x for x in [job.get("location", ""), job.get("dates", "")] if x)
                )
                meta_run.italic = True
                meta_run.font.size = Pt(9.5)
            for bullet in job.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.add_run(bullet)

    # Projects
    if data.get("projects"):
        _add_heading(doc, "Projects")
        for proj in data["projects"]:
            p = doc.add_paragraph()
            run = p.add_run(proj.get("name", ""))
            run.bold = True
            if proj.get("description"):
                doc.add_paragraph(proj["description"])
            for bullet in proj.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.add_run(bullet)

    # Education
    if data.get("education"):
        _add_heading(doc, "Education")
        for edu in data["education"]:
            p = doc.add_paragraph()
            run = p.add_run(f"{edu.get('degree', '')} — {edu.get('institution', '')}")
            run.bold = True
            meta_bits = [x for x in [edu.get("dates", ""), edu.get("details", "")] if x]
            if meta_bits:
                doc.add_paragraph(" | ".join(meta_bits))

    # Certifications
    if data.get("certifications"):
        _add_heading(doc, "Certifications")
        doc.add_paragraph(" • ".join(data["certifications"]))

    doc.save(output_path)
    return output_path


def build_cover_letter_docx(data: dict, candidate_name: str, output_path: str) -> str:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if candidate_name:
        p = doc.add_paragraph()
        run = p.add_run(candidate_name)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = ACCENT_COLOR
        doc.add_paragraph()

    if data.get("greeting"):
        doc.add_paragraph(data["greeting"])
        doc.add_paragraph()

    for para in [data.get("opening_paragraph", "")] + data.get("body_paragraphs", []) + [data.get("closing_paragraph", "")]:
        if para:
            doc.add_paragraph(para)
            doc.add_paragraph()

    if data.get("signoff"):
        doc.add_paragraph(data["signoff"])

    doc.save(output_path)
    return output_path


def build_ats_report_docx(report: dict, output_path: str) -> str:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    run = title.add_run("ATS COMPATIBILITY REPORT")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = ACCENT_COLOR

    scores = [
        ("Overall Score", report.get("overall_score")),
        ("Keyword Match", report.get("keyword_match_score")),
        ("Skills Match", report.get("skills_match_score")),
        ("Experience Match", report.get("experience_match_score")),
        ("Formatting", report.get("formatting_score")),
    ]
    for label, val in scores:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(f"{val}/100" if val is not None else "N/A")

    def bullet_section(title_text: str, items: list):
        if not items:
            return
        _add_heading(doc, title_text)
        for item in items:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(str(item))

    bullet_section("Matched Keywords", report.get("matched_keywords", []))
    bullet_section("Missing Keywords", report.get("missing_keywords", []))
    bullet_section("Formatting Issues", report.get("formatting_issues", []))
    bullet_section("Grammar Suggestions", report.get("grammar_suggestions", []))
    bullet_section("Actionable Recommendations", report.get("actionable_recommendations", []))

    doc.save(output_path)
    return output_path
